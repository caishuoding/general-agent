"""通过真实PDF/Word样本和公开HTTP行为验证文档解析、位置与限制。"""

import json
from io import BytesIO
from pathlib import Path

import httpx
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from general_agent.app import create_app
from tests.auth.support import authenticated_client
from tests.attachments.test_attachments import attachment_documents, configure, configure_settings
from tests.conversation.test_streaming import wait_message


def pdf_bytes(texts: list[str], encrypted: bool = False) -> bytes:
    """生成含真实页树和文字操作符的独立PDF样本，空字符串产生无文字页。"""
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(f'BT /F1 12 Tf 20 260 Td ({text}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = content
    if encrypted:
        writer.encrypt('fixture-password')
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def test_pdf_question_cites_actual_page_and_survives_restart(database_url: str, tmp_path: Path) -> None:
    """文本PDF经真实流式入口回答，引用到实际物理页，重启仍保留范围说明。"""
    def provider(request: httpx.Request) -> httpx.Response:
        """只用实际发给模型的页内原文生成确定性引用。"""
        payload = json.loads(request.content)
        document = attachment_documents(payload)[0]
        assert '[第 2 页]\nORBIT deadline is Friday.' in document['content']
        assert document['parse_info']['total_units'] == 2
        event = {'choices': [{'delta': {'content': 'Friday.'}, 'finish_reason': 'stop'}], 'usage': {'prompt_tokens': 30, 'completion_tokens': 5, 'total_tokens': 35}}
        return httpx.Response(200, text='data: ' + json.dumps(event) + '\n\ndata: [DONE]\n\n')

    transport = httpx.MockTransport(provider)
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, transport) as client:
        conversation = configure(client)
        result = client.post(f'/api/conversations/{conversation}/attachments', files={'files': ('报告.pdf', pdf_bytes(['Background.', 'ORBIT deadline is Friday.']), 'application/pdf')})
        assert result.status_code == 201
        attachment = client.get(f'/api/conversations/{conversation}/attachments').json()[0]
        assert attachment['status'] == 'ready'
        assert attachment['parse_info']['readable_units'] == 2
        assert '内嵌图片' in str(attachment['parse_info']['limitations'])
        assert client.post(f'/api/conversations/{conversation}/runs', json={'content': 'ORBIT deadline?'}).status_code == 202
        message = wait_message(client, conversation, 'complete')
        assert message['content'] == 'Friday.'
        assert client.get('/api/usage').json()[0]['total_tokens'] == 35
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path) as restarted:
        listed = restarted.get(f'/api/conversations/{conversation}/attachments').json()[0]
        assert listed['parse_info'] == attachment['parse_info']
        assert restarted.get(f'/api/conversations/{conversation}').json()['messages'][-1]['content'] == 'Friday.'


def docx_bytes() -> bytes:
    """生成包含空段落、中文多run和表格单元格的真实Word文档。"""
    from docx import Document

    document = Document()
    document.add_paragraph('')
    paragraph = document.add_paragraph('项目')
    paragraph.add_run('ORBIT')
    paragraph.add_run('由林女士负责。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '交付计划'
    table.cell(0, 1).text = 'CEDAR星期五交付。'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_body_and_table_have_real_structural_locations(database_url: str, tmp_path: Path) -> None:
    """DOCX正文与表格均可问答，空段落不会造成位置错号，也不虚构Word页码。"""
    def provider(request: httpx.Request) -> httpx.Response:
        """核对独立样本中的正文和表格事实确实进入外部材料及可信结构位置。"""
        payload = json.loads(request.content)
        document = attachment_documents(payload)[0]
        assert '[正文段落 2]\n项目ORBIT由林女士负责。' in document['content']
        assert 'CEDAR星期五交付。' in document['content']
        return httpx.Response(200, json={'choices': [{'message': {'content': '林女士负责，周五交付。'}}], 'usage': {'prompt_tokens': 20, 'completion_tokens': 8, 'total_tokens': 28}})

    configure_settings(database_url, tmp_path)

    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(f'/api/conversations/{conversation}/attachments', files={'files': ('计划.docx', docx_bytes())})
        assert uploaded.status_code == 201
        attachment = client.get(f'/api/conversations/{conversation}/attachments').json()[0]
        assert attachment['status'] == 'ready'
        assert attachment['parse_info']['format'] == 'docx'
        reply = client.post(f'/api/conversations/{conversation}/messages', json={'content': 'ORBIT负责人和CEDAR交付时间？'})
        assert reply.status_code == 200
        assert reply.json()['messages'][-1]['content'] == '林女士负责，周五交付。'
        assert client.get('/api/usage').json()[0]['total_tokens'] == 28


def test_compressed_package_and_extracted_text_limits_fail_truthfully(database_url: str, tmp_path: Path) -> None:
    """小压缩包不能绕过展开量限制，超大文字提取也不被当作完整理解。"""
    from zipfile import ZIP_DEFLATED, ZipFile

    package = BytesIO()
    with ZipFile(BytesIO(docx_bytes())) as original, ZipFile(package, 'w', ZIP_DEFLATED) as output:
        for name in original.namelist():
            output.writestr(name, original.read(name))
        output.writestr('word/oversized.xml', b'X' * 10_000_001)
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path) as client:
        conversation = configure(client)
        for name, content in [('compressed.docx', package.getvalue()), ('too-much-text.pdf', pdf_bytes(['X' * 2_000_001]))]:
            response = client.post(f'/api/conversations/{conversation}/attachments', files={'files': (name, content)})
            assert response.status_code == 201
            rows = client.get(f'/api/conversations/{conversation}/attachments').json()
            attachment = next(item for item in rows if item['id'] == response.json()[0]['id'])
            assert attachment['status'] == 'failed'
            assert '限制' in attachment['error']
        assert client.get('/api/usage').json() == []


def test_pdf_199_200_and_201_page_boundaries(database_url: str, tmp_path: Path) -> None:
    """真实页树199和200页可问文末，201页先拒绝且不发回答请求。"""
    requests = 0

    def provider(request: httpx.Request) -> httpx.Response:
        """从实际提供材料中返回本轮请求页的标记，检查页码不是模型自报。"""
        nonlocal requests
        payload = json.loads(request.content)
        requests += 1
        expected = 199 if requests == 1 else 200
        documents = attachment_documents(payload)
        assert any(f'PAGE_{expected}_SENTINEL' in item['content'] for item in documents)
        return httpx.Response(200, json={'choices': [{'message': {'content': '找到了末页。'}}], 'usage': {'prompt_tokens': 10, 'completion_tokens': 4, 'total_tokens': 14}})

    configure_settings(database_url, tmp_path)

    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        for count in (199, 200, 201):
            response = client.post(f'/api/conversations/{conversation}/attachments', files={'files': (f'{count}.pdf', pdf_bytes([f'PAGE_{number}_SENTINEL' for number in range(1, count + 1)]))})
            assert response.status_code == 201
            attachment = next(row for row in client.get(f'/api/conversations/{conversation}/attachments').json() if row['id'] == response.json()[0]['id'])
            assert attachment['parse_info']['total_units'] == count
            if count == 201:
                assert attachment['status'] == 'failed'
                assert '最多 200 页' in attachment['error']
                assert attachment['parse_info']['readable_units'] == 0
            else:
                assert attachment['status'] == 'ready'
                reply = client.post(f'/api/conversations/{conversation}/messages', json={'content': f'PAGE_{count}_SENTINEL?'})
                assert reply.status_code == 200
                assert reply.json()['messages'][-1]['content'] == '找到了末页。'
        assert requests == 2
        assert len(client.get('/api/usage').json()) == 2


def image_pdf_bytes(mixed: bool = False) -> bytes:
    """生成真实内嵌位图页，混合模式额外加一页可提取文字。"""
    from pypdf import PdfReader
    from pypdf.generic import NumberObject

    writer = PdfWriter()
    if mixed:
        writer.append(PdfReader(BytesIO(pdf_bytes(['Readable report on page one.']))))
    page = writer.add_blank_page(width=300, height=300)
    picture = DecodedStreamObject()
    picture.set_data(bytes([0, 128, 255, 64]))
    picture.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'), NameObject('/Width'): NumberObject(2), NameObject('/Height'): NumberObject(2), NameObject('/ColorSpace'): NameObject('/DeviceGray'), NameObject('/BitsPerComponent'): NumberObject(8)})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'): DictionaryObject({NameObject('/Im1'): picture})})
    stream = DecodedStreamObject()
    stream.set_data(b'q 100 0 0 100 20 20 cm /Im1 Do Q')
    page[NameObject('/Contents')] = stream
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def test_pdf_scans_gaps_encryption_and_corruption_are_visible(database_url: str, tmp_path: Path) -> None:
    """扫描、空白、加密、损坏均明确失败，混合PDF只声明实际可读的一页。"""
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path) as client:
        conversation = configure(client)
        samples = [('scan.pdf', image_pdf_bytes(), 'OCR'), ('blank.pdf', pdf_bytes(['']), '没有可提取文字'), ('encrypted.pdf', pdf_bytes(['Secret'], encrypted=True), '已加密'), ('broken.pdf', b'%PDF-1.7\ninvalid', '损坏'), ('fake.pdf', b'plain text', '有效 PDF')]
        for filename, content, expected in samples:
            response = client.post(f'/api/conversations/{conversation}/attachments', files={'files': (filename, content)})
            assert response.status_code == 201
            attachment = next(row for row in client.get(f'/api/conversations/{conversation}/attachments').json() if row['id'] == response.json()[0]['id'])
            assert attachment['status'] == 'failed'
            assert expected in attachment['error']
        mixed = client.post(f'/api/conversations/{conversation}/attachments', files={'files': ('mixed.pdf', image_pdf_bytes(mixed=True))}).json()[0]
        row = next(item for item in client.get(f'/api/conversations/{conversation}/attachments').json() if item['id'] == mixed['id'])
        assert row['status'] == 'ready'
        assert row['parse_info']['total_units'] == 2
        assert row['parse_info']['readable_units'] == 1
        assert row['parse_info']['coverage'] == 'partial'
        assert '第 2 页' in str(row['parse_info']['missing_units'])
        assert client.get('/api/usage').json() == []


def test_docx_revisions_missing_parts_and_invalid_packages(database_url: str, tmp_path: Path) -> None:
    """Word修订和页眉等缺口可见，伪造ZIP、损坏和空Word不能成为可读资料。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from zipfile import ZIP_DEFLATED, ZipFile

    document = Document()
    document.add_paragraph('Visible ordinary text.')
    document.sections[0].header.paragraphs[0].text = 'Header outside supported body.'
    revision = OxmlElement('w:ins')
    paragraph = OxmlElement('w:p')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'Hidden tracked insertion.'
    run.append(text)
    paragraph.append(run)
    revision.append(paragraph)
    document.element.body.append(revision)
    revised = BytesIO()
    document.save(revised)
    empty = BytesIO()
    Document().save(empty)
    fake = BytesIO()
    with ZipFile(fake, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('note.txt', 'not Word')
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path) as client:
        conversation = configure(client)
        uploaded = client.post(f'/api/conversations/{conversation}/attachments', files={'files': ('revisions.docx', revised.getvalue())}).json()[0]
        row = next(item for item in client.get(f'/api/conversations/{conversation}/attachments').json() if item['id'] == uploaded['id'])
        assert row['status'] == 'ready'
        assert row['parse_info']['coverage'] == 'partial'
        assert '修订标记' in str(row['parse_info']['missing_units'])
        assert '页眉' in str(row['parse_info']['missing_units'])
        for filename, body, error in [('fake.docx', fake.getvalue(), '有效 Word'), ('truncated.docx', docx_bytes()[:100], '损坏'), ('empty.docx', empty.getvalue(), '没有可提取')]:
            result = client.post(f'/api/conversations/{conversation}/attachments', files={'files': (filename, body)})
            assert result.status_code == 201
            row = next(item for item in client.get(f'/api/conversations/{conversation}/attachments').json() if item['id'] == result.json()[0]['id'])
            assert row['status'] == 'failed'
            assert error in row['error']
        for extension in ('.doc', '.docm'):
            assert client.post(f'/api/conversations/{conversation}/attachments', files={'files': ('wrong' + extension, docx_bytes())}).status_code == 422
        assert client.get('/api/usage').json() == []


def test_docx_entity_declarations_are_rejected_across_xml_encodings(database_url: str, tmp_path: Path) -> None:
    """UTF-8和UTF-16编码不能绕过文档XML实体声明限制。"""
    from zipfile import ZIP_DEFLATED, ZipFile

    configure_settings(database_url, tmp_path)

    with authenticated_client(database_url, tmp_path) as client:
        conversation = configure(client)
        for encoding in ('UTF-8', 'UTF-16'):
            package = BytesIO()
            with ZipFile(BytesIO(docx_bytes())) as original, ZipFile(package, 'w', ZIP_DEFLATED) as output:
                for name in original.namelist():
                    body = original.read(name)
                    if name == 'word/document.xml':
                        xml = body.decode('utf-8').replace("encoding='UTF-8'", f"encoding='{encoding}'")
                        xml = xml.replace('?>', '?><!DOCTYPE document [<!ENTITY test "test">]>', 1)
                        body = xml.encode(encoding)
                    output.writestr(name, body)
            uploaded = client.post(f'/api/conversations/{conversation}/attachments', files={'files': (f'{encoding}.docx', package.getvalue())}).json()[0]
            row = next(item for item in client.get(f'/api/conversations/{conversation}/attachments').json() if item['id'] == uploaded['id'])
            assert row['status'] == 'failed'
            assert 'XML 实体' in row['error']
        assert client.get('/api/usage').json() == []
