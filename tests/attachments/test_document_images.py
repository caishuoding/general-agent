"""验证 Word 内嵌图片生成描述切片并绑定原图身份。"""

import json
from io import BytesIO
from pathlib import Path
from time import sleep, time
from typing import Any
from uuid import UUID

import httpx
from fastapi.testclient import TestClient

from docx import Document
from PIL import Image

from general_agent.attachments.service import AttachmentService
from general_agent.infrastructure.database import Database
from general_agent.infrastructure.tokens import TokenCounter
from tests.attachments.test_attachments import configure, configure_settings
from tests.attachments.test_documents import image_pdf_bytes
from tests.attachments.test_knowledge import owner
from tests.auth.support import authenticated_client
from tests.conversation.test_streaming import wait_message


def chart_png() -> bytes:
    """生成不含文字信息的橙色图表图片，确保搜索结果只能来自视觉描述。"""
    output = BytesIO()
    Image.new('RGB', (64, 48), 'orange').save(output, 'PNG')
    return output.getvalue()


def word_with_chart() -> bytes:
    """生成带章节、独立图注和内嵌图片的真实 Word 文档。"""
    document = Document()
    document.add_heading('销售分析', level=1)
    document.add_paragraph('图 03 订单状态图及风险冻结旁路')
    document.add_picture(BytesIO(chart_png()))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def word_with_repeated_chart() -> bytes:
    """生成重复引用同一张图片的 Word 文档，供哈希缓存验证。"""
    document = Document()
    document.add_picture(BytesIO(chart_png()))
    document.add_picture(BytesIO(chart_png()))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def wait_ready(client: TestClient, conversation_id: str, attachment_id: str) -> dict[str, Any]:
    """等待后台解析和图片理解完成，返回最终附件状态。"""
    deadline = time() + 5
    while time() < deadline:
        rows = client.get(f'/api/conversations/{conversation_id}/attachments').json()
        row = next(item for item in rows if item['id'] == attachment_id)
        if row['status'] != 'processing':
            return dict(row)
        sleep(0.02)
    raise AssertionError(row)


def image_files_response() -> httpx.Response:
    """返回符合当前 Files API 合同的受控远程图片身份。"""
    created_at = int(time())
    return httpx.Response(200, json={
        'id': 'file-api-word-chart',
        'created_at': created_at,
        'expires_at': created_at + 30 * 24 * 60 * 60,
    })


def test_word_image_description_is_searchable_and_readable(
    database_url: str, tmp_path: Path,
) -> None:
    """模型直接返回纯文字，切片原样保存描述并在搜索和阅读时绑定同一原图。"""
    requests: list[dict[str, Any]] = []
    deleted: list[str] = []
    description = '季度销售趋势图：橙色柱状图显示季度销售额逐季增长。'

    def provider(request: httpx.Request) -> httpx.Response:
        """返回无 JSON 包装的图片描述，并记录模型实际收到的原图请求。"""
        if request.url.path == '/files':
            return image_files_response()
        if request.method == 'DELETE':
            deleted.append(request.url.path)
            return httpx.Response(200)
        payload = json.loads(request.content)
        requests.append(payload)
        assert payload['model'] == 'controlled'
        assert payload['stream'] is False
        assert any(
            isinstance(message.get('content'), list)
            and any(part.get('type') == 'file' for part in message['content'])
            for message in payload['messages']
        )
        return httpx.Response(200, json={
            'choices': [{'message': {'content': description}}],
            'usage': {'prompt_tokens': 20, 'completion_tokens': 10, 'total_tokens': 30},
        })

    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(
            f'/api/conversations/{conversation}/attachments',
            files={'files': ('销售分析.docx', word_with_chart())},
        ).json()[0]
        attachment = wait_ready(client, conversation, uploaded['id'])
        assert attachment['status'] == 'ready', attachment.get('error')
        assert attachment['parse_info']['image_count'] == 1
        assert attachment['parse_info']['vision_model'] == 'controlled'
        assert len(requests) == 1
        prompt = requests[0]['messages'][0]['content'][0]['text']
        assert '图 03 订单状态图及风险冻结旁路' in prompt

        service = AttachmentService(Database(database_url), tmp_path, httpx.MockTransport(provider))
        account = owner(database_url, conversation)
        searched = service.documents.search(
            UUID(conversation), account, {'query': '销售额逐季增长'},
        )['results']
        visual = next(item for item in searched if item['image_id'])
        assert visual['content'] == description
        image_id = visual['image_id']
        assert image_id

        read = service.documents.read(
            UUID(conversation),
            account,
            {'document_id': uploaded['id'], 'start_page': 1, 'end_page': 1},
            TokenCounter('test'),
            10000,
        )
        image_blocks = [block for block in read['content'] if block.get('image_id')]
        assert len(image_blocks) == 1
        assert image_blocks[0]['image_id'] == image_id
        with service.database.connect() as connection:
            chunks = connection.execute(
                'SELECT content,image_id FROM document_chunks WHERE attachment_id=%s',
                (UUID(uploaded['id']),),
            ).fetchall()
            stored_image = connection.execute(
                'SELECT caption,description FROM document_images WHERE attachment_id=%s',
                (UUID(uploaded['id']),),
            ).fetchone()
        assert all('图 03 订单状态图及风险冻结旁路' not in row['content']
                   for row in chunks if row['image_id'] is None)
        assert stored_image == {'caption': '图 03 订单状态图及风险冻结旁路', 'description': description}
        service.delete_file(UUID(uploaded['id']))
        assert deleted == ['/files/file-api-word-chart']
        assert not (tmp_path / 'attachments' / f"{uploaded['id']}.images").exists()
        database = Database(database_url)
        with database.connect() as connection:
            connection.execute(
                "UPDATE attachments SET expires_at=CURRENT_TIMESTAMP-INTERVAL '1 second' WHERE id=%s",
                (UUID(uploaded['id']),),
            )
        service.expire()
        with database.connect() as connection:
            images_count = connection.execute(
                'SELECT count(*) AS count FROM document_images WHERE attachment_id=%s',
                (UUID(uploaded['id']),),
            ).fetchone()
            chunks_count = connection.execute(
                'SELECT count(*) AS count FROM document_chunks WHERE attachment_id=%s',
                (UUID(uploaded['id']),),
            ).fetchone()
            assert images_count is not None and images_count['count'] == 0
            assert chunks_count is not None and chunks_count['count'] == 0


def test_recalled_word_image_is_sent_back_to_the_model_once(
    database_url: str, tmp_path: Path,
) -> None:
    """图片切片召回后，下一次模型请求收到原图身份并保留工具调用关系。"""
    uploads = 0
    image_requests = 0

    def provider(request: httpx.Request) -> httpx.Response:
        """先返回纯文字描述，再模拟搜索并核验下一轮原图注入。"""
        nonlocal uploads, image_requests
        if request.url.path == '/files':
            uploads += 1
            return image_files_response()
        payload = json.loads(request.content)
        if payload['stream'] is False:
            description = '季度销售趋势图：橙色柱状图显示季度销售额逐季增长。'
            return httpx.Response(200, json={
                'choices': [{'message': {'content': description}}],
            })
        has_tool = any(message.get('role') == 'tool' for message in payload['messages'])
        if not has_tool:
            return httpx.Response(200, json={'choices': [{'message': {'content': None, 'tool_calls': [{
                'id': 'image-search',
                'type': 'function',
                'function': {
                    'name': 'search_documents',
                    'arguments': json.dumps({'query': '季度销售趋势'}),
                },
            }]}}]})
        image_parts = [
            part
            for message in payload['messages']
            if isinstance(message.get('content'), list)
            for part in message['content']
            if part.get('type') == 'file'
        ]
        assert image_parts == [{'type': 'file', 'file_id': 'file-api-word-chart'}]
        image_requests += 1
        event = {
            'choices': [{'delta': {'content': '图中销售额逐季增长。'}, 'finish_reason': 'stop'}],
            'usage': {'prompt_tokens': 40, 'completion_tokens': 8, 'total_tokens': 48},
        }
        return httpx.Response(200, text='data: ' + json.dumps(event) + '\n\ndata: [DONE]\n\n')

    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(
            f'/api/conversations/{conversation}/attachments',
            files={'files': ('销售分析.docx', word_with_chart())},
        ).json()[0]
        assert wait_ready(client, conversation, uploaded['id'])['status'] == 'ready'
        response = client.post(
            f'/api/conversations/{conversation}/runs',
            json={'content': '图中的销售额趋势是什么？'},
        )
        assert response.status_code == 202
        assert wait_message(client, conversation, 'complete')['content'] == '图中销售额逐季增长。'
        assert image_requests == 1
        assert uploads == 1


def test_pdf_without_text_layer_uses_embedded_image_description(
    database_url: str, tmp_path: Path,
) -> None:
    """扫描 PDF 的纯文字图片描述可直接召回，不依赖文字层或结构化字段。"""
    def provider(request: httpx.Request) -> httpx.Response:
        """响应 PDF 图片上传，并让模型直接输出一段自然语言描述。"""
        if request.url.path == '/files':
            return httpx.Response(200, json={
                'id': 'file-api-pdf-scan',
                'created_at': int(time()),
                'expires_at': int(time()) + 30 * 24 * 60 * 60,
            })
        description = '蓝色扫描画面：蓝灰色方块组成的扫描页图像。'
        return httpx.Response(200, json={
            'choices': [{'message': {'content': description}}],
        })

    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(
            f'/api/conversations/{conversation}/attachments',
            files={'files': ('扫描.pdf', image_pdf_bytes())},
        ).json()[0]
        attachment = wait_ready(client, conversation, uploaded['id'])
        assert attachment['status'] == 'ready', attachment.get('error')
        assert attachment['parse_info']['image_count'] == 1

        service = AttachmentService(Database(database_url), tmp_path)
        searched = service.documents.search(
            UUID(conversation),
            owner(database_url, conversation),
            {'query': '蓝色扫描画面'},
        )['results']
        visual = next(item for item in searched if item['image_id'])
        assert '扫描页图像' in visual['content']


def test_word_text_remains_available_when_image_description_fails(
    database_url: str, tmp_path: Path,
) -> None:
    """视觉服务失败只标记图片，不能把仍然可读的 Word 正文判为失败。"""
    def provider(request: httpx.Request) -> httpx.Response:
        """让图片上传失败，并确认应用不会继续调用聊天模型生成描述。"""
        assert request.url.path == '/files'
        return httpx.Response(503)

    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(
            f'/api/conversations/{conversation}/attachments',
            files={'files': ('销售分析.docx', word_with_chart())},
        ).json()[0]
        attachment = wait_ready(client, conversation, uploaded['id'])
        assert attachment['status'] == 'ready'
        assert attachment['parse_info']['image_count'] == 1
        assert attachment['parse_info']['image_described_count'] == 0
        assert attachment['parse_info']['image_failed_count'] == 1

        service = AttachmentService(Database(database_url), tmp_path)
        searched = service.documents.search(
            UUID(conversation),
            owner(database_url, conversation),
            {'query': '销售分析'},
        )['results']
        assert any('销售分析' in item['content'] for item in searched)
        assert all('图 03 订单状态图及风险冻结旁路' not in item['content'] for item in searched)
        assert all(item['image_id'] is None for item in searched)


def test_duplicate_word_images_reuse_one_visual_description(
    database_url: str, tmp_path: Path,
) -> None:
    """同一文档内字节相同的图片只上传和描述一次，但保留两个出现位置。"""
    uploads = 0
    descriptions = 0

    def provider(request: httpx.Request) -> httpx.Response:
        """统计图片上传及纯文字描述的真实请求次数。"""
        nonlocal uploads, descriptions
        if request.url.path == '/files':
            uploads += 1
            return image_files_response()
        descriptions += 1
        description = '重复图表：橙色图表示例。'
        return httpx.Response(200, json={
            'choices': [{'message': {'content': description}}],
        })

    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path, httpx.MockTransport(provider)) as client:
        conversation = configure(client)
        uploaded = client.post(
            f'/api/conversations/{conversation}/attachments',
            files={'files': ('重复图片.docx', word_with_repeated_chart())},
        ).json()[0]
        attachment = wait_ready(client, conversation, uploaded['id'])
        assert attachment['status'] == 'ready'
        assert attachment['parse_info']['image_count'] == 2
        assert attachment['parse_info']['image_described_count'] == 2
        assert uploads == 1
        assert descriptions == 1
