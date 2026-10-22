"""验证 Word 原生结构、数字保真、页码映射及切片边界。"""

from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from general_agent.attachments.documents import ParsedDocument, parse_docx, parse_document
from general_agent.attachments.knowledge import ChunkPolicy
from general_agent.attachments.service import AttachmentService
from general_agent.attachments import service as service_module
from general_agent.attachments.tools import DocumentTools
from general_agent.attachments.pagination import paginate_word
from general_agent.infrastructure.database import Database
from general_agent.infrastructure.tokens import TokenCounter
from tests.auth.support import authenticated_client
from tests.attachments.test_attachments import configure, configure_settings
from tests.attachments.test_knowledge import owner


def structured_word() -> bytes:
    """构造数字、分节标题、表格及明确分页均可校验的独立 Word 样本。"""
    doc = Document()
    doc.add_heading('1 资金与成果', level=1)
    doc.add_paragraph('WordStream 样本13474活动，中位CPC5.42美元。')
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['资金', '预算', '负责人']):
        cell.text = text
    for i in range(20):
        for cell, text in zip(table.add_row().cells, [f'项目{i:02d}', f'{400+i}万元', f'负责人{i:02d}']):
            cell.text = text
    heading = doc.add_heading('2 会员定价', level=1)
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph('CPC5.42美元，手机号1933，条款14至17条。')
    paragraph = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to('https://example.com/rates', RT.HYPERLINK, is_external=True))
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = '官方计价 CPC5.42美元'
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def test_docx_parser_keeps_original_headings_rows_and_digits(tmp_path: Path) -> None:
    """不经过 PDF 的原生解析必须保留标题类型、完整表格行及精确数字。"""
    path = tmp_path / 'native.docx'
    path.write_bytes(structured_word())
    parsed = parse_docx(path)
    assert parsed.error is None
    assert parsed.units
    assert any(unit.kind == 'heading' and unit.content.startswith('1 资金') for unit in parsed.units)
    rows = [unit for unit in parsed.units if unit.kind == 'table_row']
    assert len(rows) == 21
    assert rows[0].content == '| 资金 | 预算 | 负责人 |\n| --- | --- | --- |\n'
    assert rows[1].content == '| 项目00 | 400万元 | 负责人00 |\n'
    assert 'CPC5.42' in ''.join(unit.content for unit in parsed.units)


def test_word_upload_indexes_native_text_and_reads_bookmarked_pages(
    database_url: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """真实上传不得回读 PDF 正文；原生数字、Markdown 表格和分页映射需一直保留到工具结果。"""
    original_parse = parse_document

    def native_only(path: Path, extension: str) -> ParsedDocument:
        """确保 Word 路径只解析原件，禁止用 PDF 提取文字覆盖原生结构。"""
        assert extension == '.docx', 'Word 不得从分页 PDF 再次抽取正文。'
        return original_parse(path, extension)

    monkeypatch.setattr(service_module, 'parse_document', native_only)
    configure_settings(database_url, tmp_path)
    with authenticated_client(database_url, tmp_path) as client:
        conversation = configure(client)
        identifier = client.post(f'/api/conversations/{conversation}/attachments',
                                 files={'files': ('native.docx', structured_word())}).json()[0]['id']
        listed = client.get(f'/api/conversations/{conversation}/attachments').json()[0]
        assert listed['status'] == 'ready', listed['error']
        database = Database(database_url)
        tools = DocumentTools(database, ChunkPolicy())
        account = owner(database_url, conversation)
        result = tools.read(UUID(conversation), account, {'document_id': identifier, 'start_page': 2, 'end_page': 2},
                            TokenCounter('test'), 10000)
        text = ''.join(block['content'] for block in result['content'])
        assert 'CPC5.42美元，手机号1933，条款14至17条。' in text
        assert '官方计价 CPC5.42美元' in text
        assert '项目00' not in text
        with database.connect() as connection:
            chunks = connection.execute('SELECT * FROM document_chunks WHERE attachment_id=%s ORDER BY ordinal',
                                        (UUID(identifier),)).fetchall()
            assert all(TokenCounter('test').count(chunk['content']) <= 1000 for chunk in chunks)
            assert all(chunk['token_count'] == TokenCounter('test').count(chunk['content']) for chunk in chunks)
            assert all(not ('项目' in chunk['content'] and '会员定价' in chunk['content']) for chunk in chunks)
            assert sum(chunk['content'].count('项目00') for chunk in chunks) == 1
        attachments = AttachmentService(database, tmp_path)
        before = listed['expires_at']
        attachments.recover()
        assert client.get(f'/api/conversations/{conversation}/attachments').json()[0]['expires_at'] == before


def test_word_bookmarks_cover_a_paragraph_across_multiple_pages(tmp_path: Path) -> None:
    """单一长段落跨页时书签必须覆盖真实起止页，不能用段首页冒充整段范围。"""
    doc = Document()
    doc.add_paragraph('LONGWORD ' * 1200)
    path = tmp_path / 'long.docx'
    doc.save(str(path))
    count, positions = paginate_word(path, tmp_path / 'long.pdf')
    assert count > 1
    assert positions == {1: (1, count)}
