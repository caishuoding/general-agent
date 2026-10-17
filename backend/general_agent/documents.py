"""提取文档中可读取的文字及真实位置，保留没有覆盖的内容说明。"""

import ctypes
import os
import subprocess
import sys
from ctypes import Structure
from ctypes import wintypes
from tempfile import TemporaryFile
from time import monotonic, sleep
from pathlib import Path
from typing import NamedTuple
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as WordDocument
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from pydantic import BaseModel, Field
from pypdf import PdfReader, filters

from general_agent.document_limits import (
    MAX_DOCUMENT_CPU_SECONDS,
    MAX_DOCUMENT_OUTPUT_BYTES,
    MAX_DOCUMENT_RSS_BYTES,
)

MAX_TEXT_CHARACTERS = 2_000_000


class _FileTime(Structure):
    """描述 Windows 进程时间计数器的两个 32 位部分。"""

    _fields_ = [('low', wintypes.DWORD), ('high', wintypes.DWORD)]


class _ProcessMemoryCounters(Structure):
    """接收 Windows GetProcessMemoryInfo 返回的工作集统计。"""

    _fields_ = [
        ('cb', wintypes.DWORD), ('page_fault_count', wintypes.DWORD),
        ('peak_working_set_size', ctypes.c_size_t), ('working_set_size', ctypes.c_size_t),
        ('quota_peak_paged_pool_usage', ctypes.c_size_t), ('quota_paged_pool_usage', ctypes.c_size_t),
        ('quota_peak_non_paged_pool_usage', ctypes.c_size_t), ('quota_non_paged_pool_usage', ctypes.c_size_t),
        ('pagefile_usage', ctypes.c_size_t), ('peak_pagefile_usage', ctypes.c_size_t),
    ]


class _ProcessUsage(NamedTuple):
    """保存文档工作进程的工作集字节和累计 CPU 秒数。"""

    rss_bytes: int
    cpu_seconds: float


def _file_time_seconds(value: _FileTime) -> float:
    """将 Windows 的 100 纳秒 FILETIME 计数转换为秒。"""
    return float((int(value.high) << 32) | int(value.low)) / 10_000_000


def _windows_process_usage(process_id: int) -> _ProcessUsage | None:
    """读取 Windows 子进程的工作集字节和累计 CPU 秒数。"""
    loader = getattr(ctypes, 'WinDLL', None)
    if loader is None:
        return None
    kernel32 = loader('kernel32', use_last_error=True)
    psapi = loader('psapi', use_last_error=True)
    kernel32.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel32.OpenProcess.restype = wintypes.HANDLE
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    kernel32.GetProcessTimes.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(_FileTime), ctypes.POINTER(_FileTime),
        ctypes.POINTER(_FileTime), ctypes.POINTER(_FileTime),
    ]
    kernel32.GetProcessTimes.restype = wintypes.BOOL
    psapi.GetProcessMemoryInfo.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(_ProcessMemoryCounters), wintypes.DWORD,
    ]
    psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
    handle = kernel32.OpenProcess(0x1000, False, process_id)
    if not handle:
        return None
    try:
        counters = _ProcessMemoryCounters()
        counters.cb = ctypes.sizeof(counters)
        created, exited, kernel, user = _FileTime(), _FileTime(), _FileTime(), _FileTime()
        if not psapi.GetProcessMemoryInfo(handle, ctypes.byref(counters), counters.cb):
            return None
        if not kernel32.GetProcessTimes(handle, ctypes.byref(created), ctypes.byref(exited), ctypes.byref(kernel), ctypes.byref(user)):
            return None
        return _ProcessUsage(int(counters.working_set_size), _file_time_seconds(kernel) + _file_time_seconds(user))
    finally:
        kernel32.CloseHandle(handle)


class DocumentInfo(BaseModel):
    """公开的文字提取范围，不代表模型已经理解整份文档。"""
    format: str
    total_units: int = 0
    readable_units: int = 0
    missing_units: list[str] = Field(default_factory=list)
    limitations: list[str] = Field(default_factory=list)
    coverage: str = 'complete'


class ParsedDocument(BaseModel):
    """可保存的文字片段、处理范围及真实失败原因。"""
    chunks: list[tuple[str, str]] = Field(default_factory=list)
    info: DocumentInfo
    error: str | None = None


def split_unit(location: str, text: str) -> list[tuple[str, str]]:
    """短单元保留直接位置，长单元分片时补充原单元内字符范围。"""
    if len(text) <= 1500:
        return [(location, text)] if text.strip() else []
    return [(f'{location}，字符 {offset + 1}–{min(offset + 1500, len(text))}', text[offset:offset + 1500]) for offset in range(0, len(text), 1500)]


def parse_pdf(path: Path) -> ParsedDocument:
    """按物理页序限额读取文字，空页或错误页作为缺口保留，不做OCR。"""
    info = DocumentInfo(format='pdf', limitations=['仅提取文字；复杂图表、内嵌图片和特殊排版未被理解，需要将相关图像单独上传。'])
    try:
        filters.ZLIB_MAX_OUTPUT_LENGTH = 10_000_000
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            return ParsedDocument(info=info, error='PDF 已加密，请上传未加密副本。')
        info.total_units = len(reader.pages)
        if info.total_units > 200:
            return ParsedDocument(info=info, error='单份 PDF 最多 200 页，请拆分后上传。')
        chunks: list[tuple[str, str]] = []
        text_size = 0
        for number, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text()
            except Exception:
                info.missing_units.append(f'第 {number} 页提取失败')
                continue
            text_size += len(text)
            if text_size > MAX_TEXT_CHARACTERS:
                return ParsedDocument(info=info, error="文档提取文字超过 2,000,000 字符资源限制，请拆分后上传。")
            if not text.strip():
                info.missing_units.append(f'第 {number} 页没有可提取文字（可能是空白或扫描图像）')
            else:
                info.readable_units += 1
                chunks.extend(split_unit(f'第 {number} 页', text))
        info.coverage = 'partial' if info.missing_units else 'complete'
        if not chunks:
            return ParsedDocument(info=info, error='PDF 没有可提取文字；可能是扫描版或空白文档，本版本不提供 OCR。')
        return ParsedDocument(chunks=chunks, info=info)
    except Exception:
        return ParsedDocument(info=info, error='PDF 损坏或不是有效 PDF，无法读取文字。')


def word_units(container: WordDocument | _Cell, prefix: str = '', depth: int = 0) -> list[tuple[str, str]]:
    """按真实正文顺序展开段落和表格，保留空段落计数且不生成分页信息。"""
    if depth > 8:
        raise ValueError('DOCX 表格嵌套过深，请简化文档后上传。')
    units: list[tuple[str, str]] = []
    paragraphs, tables = 0, 0
    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraphs += 1
            location = f'{prefix}，段落 {paragraphs}' if prefix else f'正文段落 {paragraphs}'
            units.append((location, item.text))
        else:
            tables += 1
            table_prefix = f'{prefix}，嵌套表格 {tables}' if prefix else f'表格 {tables}'
            seen: set[object] = set()
            for row_number, row in enumerate(item.rows, 1):
                for column_number, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    units.extend(word_units(cell, f'{table_prefix}，第 {row_number} 行，第 {column_number} 列', depth + 1))
    return units


def parse_docx(path: Path) -> ParsedDocument:
    """先限制ZIP展开与XML实体，再解析正文表格并说明修订及附属内容缺口。"""
    info = DocumentInfo(format='docx', limitations=['仅提取正文及表格文字；复杂图表、内嵌图片和特殊排版未被理解，需要将相关图像单独上传。'])
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 50_000_000 or any(item.file_size > 10_000_000 for item in entries):
                raise ValueError('DOCX 压缩包展开量超过资源限制（2000 个条目、单项 10 MB、总计 50 MB），请拆分或简化文档。')
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise ValueError('DOCX 压缩包存在重复结构，不能可靠读取。')
            for name in names:
                if name.endswith(('.xml', '.rels')):
                    xml = archive.read(name)
                    declaration = xml.replace(b'\x00', b'').upper()
                    if b'<!DOCTYPE' in declaration or b'<!ENTITY' in declaration:
                        raise ValueError('DOCX 包含不支持的 XML 实体声明，已拒绝解析。')
            content_types = ElementTree.fromstring(archive.read('[Content_Types].xml'))
            if not any(item.attrib.get('PartName') == '/word/document.xml' and item.attrib.get('ContentType') == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml' for item in content_types):
                raise ValueError('文件不是支持的 DOCX 文档，请勿用 DOC、DOCM 或普通 ZIP 改扩展名。')
            root = ElementTree.fromstring(archive.read('word/document.xml'))
            tags = {item.tag.rsplit('}', 1)[-1] for item in root.iter()}
            if tags & {'ins', 'del', 'moveFrom', 'moveTo'}:
                info.missing_units.append('修订标记中的文字未提取；请接受或拒绝修订后重新上传。')
            if 'txbxContent' in tags:
                info.missing_units.append('文本框文字不在当前提取范围。')
            if any(name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes', 'word/comments')) for name in names):
                info.missing_units.append('页眉、页脚、脚注、尾注及批注不在当前提取范围。')
        units = word_units(Document(str(path)))
        if sum(len(text) for _, text in units) > MAX_TEXT_CHARACTERS:
            raise ValueError('文档提取文字超过 2,000,000 字符资源限制，请拆分后上传。')
        info.total_units = len(units)
        info.readable_units = sum(bool(text.strip()) for _, text in units)
        info.coverage = 'partial' if info.missing_units else 'complete'
        chunks = [chunk for location, text in units for chunk in split_unit(location, text)]
        if not chunks:
            return ParsedDocument(info=info, error='DOCX 没有可提取的正文或表格文字；图像和扫描内容不会自动识别。')
        return ParsedDocument(chunks=chunks, info=info)
    except ValueError as error:
        return ParsedDocument(info=info, error=str(error))
    except Exception:
        return ParsedDocument(info=info, error='DOCX 损坏、受保护或不是有效 Word 文档，无法读取文字。')


def parse_document(path: Path, extension: str) -> ParsedDocument:
    """隔离解析并监控实际内存及墙钟时间，异常子进程不能拖住应用线程或回显原始错误。"""
    info = DocumentInfo(format=extension.lstrip('.'))
    environment = {
        'PYTHONPATH': str(Path(__file__).resolve().parents[1]),
        'PATH': os.environ.get('PATH', ''),
        'LANG': 'en_US.UTF-8',
        'PYTHONIOENCODING': 'utf-8',
        'PYTHONUTF8': '1',
    }
    with TemporaryFile() as output:
        process = subprocess.Popen([sys.executable, '-m', 'general_agent.document_worker', str(path.resolve()), extension], stdout=output, stderr=subprocess.DEVNULL, env=environment)
        try:
            deadline = monotonic() + 20
            while process.poll() is None:
                if monotonic() >= deadline:
                    return ParsedDocument(info=info, error='文档处理超过 20 秒资源限制，请拆分或简化后上传。')
                if sys.platform == 'win32':
                    usage = _windows_process_usage(process.pid)
                    if usage is not None and usage.rss_bytes > MAX_DOCUMENT_RSS_BYTES:
                        return ParsedDocument(info=info, error='文档解析超过 512 MB 内存资源限制，请拆分或简化后上传。')
                    if usage is not None and usage.cpu_seconds > MAX_DOCUMENT_CPU_SECONDS:
                        return ParsedDocument(info=info, error='文档解析超过 CPU 资源限制，请拆分或简化后上传。')
                else:
                    observed = subprocess.run(['/bin/ps', '-o', 'rss=', '-p', str(process.pid)], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=2, check=False)
                    if observed.returncode == 0 and int(observed.stdout.strip() or b'0') * 1024 > MAX_DOCUMENT_RSS_BYTES:
                        return ParsedDocument(info=info, error='文档解析超过 512 MB 内存资源限制，请拆分或简化后上传。')
                sleep(0.025)
            if process.returncode != 0 or output.tell() > MAX_DOCUMENT_OUTPUT_BYTES:
                return ParsedDocument(info=info, error='文档解析失败或超过CPU、内存、输出资源限制，请拆分或简化后上传。')
            output.seek(0)
            return ParsedDocument.model_validate_json(output.read(16_000_001))
        except Exception:
            return ParsedDocument(info=info, error='受限文档解析未完成，请检查文件后重新上传。')
        finally:
            if process.poll() is None:
                process.kill()
            process.wait(timeout=3)
