"""提取文档中可读取的文字及真实位置，保留没有覆盖的内容说明。"""

import ctypes
import os
import subprocess
import sys
from ctypes import Structure
from ctypes import wintypes
from tempfile import TemporaryFile
from time import monotonic, sleep
from pathlib import Path
from typing import NamedTuple
from types import SimpleNamespace
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document

from pydantic import BaseModel, Field
from pypdf import PdfReader, filters
import pdfplumber

from general_agent.attachments.limits import (
    MAX_DOCUMENT_CPU_SECONDS,
    MAX_DOCUMENT_OUTPUT_BYTES,
    MAX_DOCUMENT_RSS_BYTES,
)
from general_agent.attachments.knowledge import DocumentUnit, text_units, MAX_DOCUMENT_UNITS
from general_agent.attachments.word_structure import word_units
from general_agent.attachments.pdf_structure import page_units
from general_agent.attachments.pdf_tables import annotate_cross_page_tables
from general_agent.attachments.pdf_text import PDF_TEXT_SOURCE, document_layout, merge_page_paragraphs
from general_agent.attachments.document_images import (
    ParsedImage,
    extract_pdf_images,
    extract_word_images,
    figure_caption,
)

MAX_TEXT_CHARACTERS = 2_000_000


class _FileTime(Structure):
    """描述 Windows 进程时间计数器的两个 32 位部分。"""

    _fields_ = [('low', wintypes.DWORD), ('high', wintypes.DWORD)]


class _ProcessMemoryCounters(Structure):
    """接收 Windows GetProcessMemoryInfo 返回的工作集统计。"""

    _fields_ = [
        ('cb', wintypes.DWORD), ('page_fault_count', wintypes.DWORD),
        ('peak_working_set_size', ctypes.c_size_t), ('working_set_size', ctypes.c_size_t),
        ('quota_peak_paged_pool_usage', ctypes.c_size_t), ('quota_paged_pool_usage', ctypes.c_size_t),
        ('quota_peak_non_paged_pool_usage', ctypes.c_size_t), ('quota_non_paged_pool_usage', ctypes.c_size_t),
        ('pagefile_usage', ctypes.c_size_t), ('peak_pagefile_usage', ctypes.c_size_t),
    ]


class _ProcessUsage(NamedTuple):
    """保存文档工作进程的工作集字节和累计 CPU 秒数。"""

    rss_bytes: int
    cpu_seconds: float


def _file_time_seconds(value: _FileTime) -> float:
    """将 Windows 的 100 纳秒 FILETIME 计数转换为秒。"""
    return float((int(value.high) << 32) | int(value.low)) / 10_000_000


def _windows_process_usage(process_id: int) -> _ProcessUsage | None:
    """读取 Windows 子进程的工作集字节和累计 CPU 秒数。"""
    loader = getattr(ctypes, 'WinDLL', None)
    if loader is None:
        return None
    kernel32 = loader('kernel32', use_last_error=True)
    psapi = loader('psapi', use_last_error=True)
    kernel32.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel32.OpenProcess.restype = wintypes.HANDLE
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    kernel32.GetProcessTimes.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(_FileTime), ctypes.POINTER(_FileTime),
        ctypes.POINTER(_FileTime), ctypes.POINTER(_FileTime),
    ]
    kernel32.GetProcessTimes.restype = wintypes.BOOL
    psapi.GetProcessMemoryInfo.argtypes = [
        wintypes.HANDLE, ctypes.POINTER(_ProcessMemoryCounters), wintypes.DWORD,
    ]
    psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
    handle = kernel32.OpenProcess(0x1000, False, process_id)
    if not handle:
        return None
    try:
        counters = _ProcessMemoryCounters()
        counters.cb = ctypes.sizeof(counters)
        created, exited, kernel, user = _FileTime(), _FileTime(), _FileTime(), _FileTime()
        if not psapi.GetProcessMemoryInfo(handle, ctypes.byref(counters), counters.cb):
            return None
        if not kernel32.GetProcessTimes(handle, ctypes.byref(created), ctypes.byref(exited), ctypes.byref(kernel), ctypes.byref(user)):
            return None
        return _ProcessUsage(int(counters.working_set_size), _file_time_seconds(kernel) + _file_time_seconds(user))
    finally:
        kernel32.CloseHandle(handle)


class DocumentInfo(BaseModel):
    """公开的文字提取范围，不代表模型已经理解整份文档。"""
    format: str
    total_units: int = 0
    readable_units: int = 0
    missing_units: list[str] = Field(default_factory=list)
    limitations: list[str] = Field(default_factory=list)
    coverage: str = 'complete'
    text_source: str | None = None
    image_count: int = 0


class ParsedDocument(BaseModel):
    """传给切片器的临时解析块、处理范围及真实失败原因。"""
    info: DocumentInfo
    error: str | None = None
    units: list[DocumentUnit] = Field(default_factory=list)
    images: list[ParsedImage] = Field(default_factory=list)


def attach_pdf_captions(units: list[DocumentUnit], images: list[ParsedImage]) -> list[ParsedImage]:
    """把同页图注与 PDF 图片按原文顺序绑定；数量不一致时只绑定可确定的部分。"""
    captions: dict[int, list[str]] = {}
    for unit in units:
        if unit.page is None:
            continue
        caption = figure_caption(unit.content.strip())
        if caption is not None:
            captions.setdefault(unit.page, []).append(caption)
    by_page: dict[int, list[int]] = {}
    for index, image in enumerate(images):
        if image.page is not None:
            by_page.setdefault(image.page, []).append(index)
    result = list(images)
    for page, indexes in by_page.items():
        page_captions = captions.get(page, [])
        if len(indexes) == 1 and page_captions:
            result[indexes[0]] = result[indexes[0]].model_copy(update={'caption': page_captions[0]})
        elif len(indexes) == len(page_captions):
            for index, caption in zip(indexes, page_captions):
                result[index] = result[index].model_copy(update={'caption': caption})
    return result


def exclude_caption_units(units: list[DocumentUnit], images: list[ParsedImage]) -> None:
    """从普通正文单位中移除已绑定图注，保留单位数量以维持分页位置映射。"""
    captions = {image.caption.strip() for image in images if image.caption and image.caption.strip()}
    if not captions:
        return
    for unit in units:
        lines = unit.content.splitlines(keepends=True)
        remaining = [line for line in lines if figure_caption(line.strip()) not in captions]
        if len(remaining) != len(lines):
            unit.content = ''.join(remaining)


def parse_pdf(path: Path, image_directory: Path | None = None) -> ParsedDocument:
    """按单栏版面还原 PDF 标题、自然段与表格位置，图片由后续视觉阶段生成描述。"""
    info = DocumentInfo(format='pdf', text_source=PDF_TEXT_SOURCE, limitations=['面向单栏文字型 PDF；标题和段落按版面启发式还原，复杂矢量图保留原图供视觉理解但不还原编辑结构。', '表格保留真实表头、合并单元格与续表；无法可靠识别的数据矩阵使用原区域截图。'])
    try:
        filters.ZLIB_MAX_OUTPUT_LENGTH = 10_000_000
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            return ParsedDocument(info=info, error='PDF 已加密，请上传未加密副本。')
        info.total_units = len(reader.pages)
        if info.total_units > 200:
            return ParsedDocument(info=info, error='单份 PDF 最多 200 页，请拆分后上传。')
        images: list[ParsedImage] = []
        units: list[DocumentUnit] = []
        text_size, table_count = 0, 0
        with pdfplumber.open(path) as layout:
            try:
                text_layout = document_layout(layout, reader)
            except ValueError as error:
                return ParsedDocument(info=info, error=str(error))
            info.missing_units.extend(f'第 {number} 页版面统计未完成，已保留逐页文字回退'
                                      for number in sorted(text_layout.incomplete_pages))
            for number, page in enumerate(reader.pages, 1):
                try:
                    text = page.extract_text()
                except Exception:
                    info.missing_units.append(f'第 {number} 页提取失败')
                    units.extend(text_units('', number))
                    continue
                text_size += len(text)
                if text_size > MAX_TEXT_CHARACTERS:
                    return ParsedDocument(info=info, error='文档提取文字超过 2,000,000 字符资源限制，请拆分后上传。')
                structure = None
                try:
                    structure = layout.pages[number - 1]
                    page_content, count, page_pictures = page_units(structure, number, table_count, text_layout,
                                                                    image_directory, len(images))
                    table_count += count
                    images.extend(page_pictures)
                except Exception as error:
                    if isinstance(error, ValueError) and '资源限制' in str(error):
                        return ParsedDocument(info=info, error=str(error))
                    info.missing_units.append(f'第 {number} 页表格识别失败，已回退纯文字提取')
                    page_content = text_units(text, number)
                    fallback = extract_pdf_images(SimpleNamespace(pages=[page]), image_directory)
                    for image in fallback:
                        image.ordinal = len(images) + 1
                        image.page = image.end_page = image.paragraph = number
                        image.location = f'第 {number} 页内嵌图片 {image.ordinal}'
                        images.append(image)
                finally:
                    if structure is not None:
                        structure.close()
                units.extend(page_content)
                if len(units) > MAX_DOCUMENT_UNITS:
                    return ParsedDocument(info=info, error='文档段落数量超过 20,000 个资源限制，请拆分后上传。')
                if not any(unit.content.strip() for unit in page_content):
                    info.missing_units.append(f'第 {number} 页没有可提取文字（可能是空白或扫描图像）')
                else:
                    info.readable_units += 1
        if sum(len(unit.content) for unit in units) > MAX_TEXT_CHARACTERS:
            return ParsedDocument(info=info, error='文档提取文字超过 2,000,000 字符资源限制，请拆分后上传。')
        annotate_cross_page_tables(units, text_layout.page_sizes)
        units = merge_page_paragraphs(units, text_layout)
        info.image_count = len(images)
        exclude_caption_units(units, images)
        info.readable_units = len({page for unit in units if unit.content.strip()
                                   for page in range(unit.page or 1, (unit.end_page or unit.page or 1) + 1)})
        info.coverage = 'partial' if info.missing_units else 'complete'
        if not info.readable_units:
            if not images:
                return ParsedDocument(info=info, error='PDF 没有可提取文字或图片；可能是空白文档。')
            info.coverage = 'partial'
            info.missing_units.append('PDF 没有文字层，正在使用页面内嵌图片进行视觉理解')
        return ParsedDocument(info=info, units=units, images=images)
    except Exception:
        return ParsedDocument(info=info, error='PDF 损坏或不是有效 PDF，无法读取文字。')


def parse_docx(path: Path, image_directory: Path | None = None) -> ParsedDocument:
    """校验 ZIP 与 XML 后提取正文、表格及内嵌图片，并保留全部真实位置。"""
    info = DocumentInfo(format='docx', text_source='word_native_markdown_v3', limitations=['SmartArt、EMF、WMF 和复杂矢量图表暂不解析。'])
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 50_000_000 or any(item.file_size > 10_000_000 for item in entries):
                raise ValueError('DOCX 压缩包展开量超过资源限制（2000 个条目、单项 10 MB、总计 50 MB），请拆分或简化文档。')
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise ValueError('DOCX 压缩包存在重复结构，不能可靠读取。')
            for name in names:
                if name.endswith(('.xml', '.rels')):
                    xml = archive.read(name)
                    declaration = xml.replace(b'\x00', b'').upper()
                    if b'<!DOCTYPE' in declaration or b'<!ENTITY' in declaration:
                        raise ValueError('DOCX 包含不支持的 XML 实体声明，已拒绝解析。')
            content_types = ElementTree.fromstring(archive.read('[Content_Types].xml'))
            if not any(item.attrib.get('PartName') == '/word/document.xml' and item.attrib.get('ContentType') == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml' for item in content_types):
                raise ValueError('文件不是支持的 DOCX 文档，请勿用 DOC、DOCM 或普通 ZIP 改扩展名。')
            root = ElementTree.fromstring(archive.read('word/document.xml'))
            tags = {item.tag.rsplit('}', 1)[-1] for item in root.iter()}
            if tags & {'ins', 'del', 'moveFrom', 'moveTo'}:
                info.missing_units.append('修订标记中的文字未提取；请接受或拒绝修订后重新上传。')
            if 'txbxContent' in tags:
                info.missing_units.append('文本框文字不在当前提取范围。')
            if any(name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes', 'word/comments')) for name in names):
                info.missing_units.append('页眉、页脚、脚注、尾注及批注不在当前提取范围。')
        document = Document(str(path))
        units = word_units(document)
        try:
            skipped_images: list[str] = []
            images = extract_word_images(document, image_directory, skipped_images)
            info.missing_units.extend(skipped_images)
        except ValueError as error:
            return ParsedDocument(info=info, error=str(error))
        info.image_count = len(images)
        exclude_caption_units(units, images)
        if sum(len(unit.content) for unit in units) > MAX_TEXT_CHARACTERS:
            raise ValueError('文档提取文字超过 2,000,000 字符资源限制，请拆分后上传。')
        info.total_units = len(units)
        info.readable_units = sum(bool(unit.content.strip()) for unit in units)
        info.coverage = 'partial' if info.missing_units else 'complete'
        if not info.readable_units:
            if not images:
                return ParsedDocument(info=info, error='DOCX 没有可提取的正文、表格或图片。')
            info.coverage = 'partial'
            info.missing_units.append('DOCX 没有正文文字，正在使用内嵌图片进行视觉理解')
        return ParsedDocument(info=info, units=units, images=images)
    except ValueError as error:
        return ParsedDocument(info=info, error=str(error))
    except Exception:
        return ParsedDocument(info=info, error='DOCX 损坏、受保护或不是有效 Word 文档，无法读取文字。')


def parse_document(path: Path, extension: str, image_directory: Path) -> ParsedDocument:
    """隔离解析并监控实际内存及墙钟时间，异常子进程不能拖住应用线程或回显原始错误。"""
    info = DocumentInfo(format=extension.lstrip('.'))
    environment = {
        # 工作进程需要包含 general_agent 包的 backend 根目录，而不是包目录本身。
        'PYTHONPATH': str(Path(__file__).resolve().parents[2]),
        'PATH': os.environ.get('PATH', ''),
        'LANG': 'en_US.UTF-8',
        'PYTHONIOENCODING': 'utf-8',
        'PYTHONUTF8': '1',
    }
    with TemporaryFile() as output:
        process = subprocess.Popen(
            [sys.executable, '-m', 'general_agent.attachments.worker', str(path.resolve()), extension,
             str(image_directory.resolve())],
            stdout=output,
            stderr=subprocess.DEVNULL,
            env=environment,
        )
        try:
            deadline = monotonic() + 20
            while process.poll() is None:
                if monotonic() >= deadline:
                    return ParsedDocument(info=info, error='文档处理超过 20 秒资源限制，请拆分或简化后上传。')
                if sys.platform == 'win32':
                    usage = _windows_process_usage(process.pid)
                    if usage is not None and usage.rss_bytes > MAX_DOCUMENT_RSS_BYTES:
                        return ParsedDocument(info=info, error='文档解析超过 512 MB 内存资源限制，请拆分或简化后上传。')
                    if usage is not None and usage.cpu_seconds > MAX_DOCUMENT_CPU_SECONDS:
                        return ParsedDocument(info=info, error='文档解析超过 CPU 资源限制，请拆分或简化后上传。')
                else:
                    observed = subprocess.run(['/bin/ps', '-o', 'rss=', '-p', str(process.pid)], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=2, check=False)
                    if observed.returncode == 0 and int(observed.stdout.strip() or b'0') * 1024 > MAX_DOCUMENT_RSS_BYTES:
                        return ParsedDocument(info=info, error='文档解析超过 512 MB 内存资源限制，请拆分或简化后上传。')
                sleep(0.025)
            if process.returncode != 0 or output.tell() > MAX_DOCUMENT_OUTPUT_BYTES:
                return ParsedDocument(info=info, error='文档解析失败或超过CPU、内存、输出资源限制，请拆分或简化后上传。')
            output.seek(0)
            return ParsedDocument.model_validate_json(output.read(16_000_001))
        except Exception:
            return ParsedDocument(info=info, error='受限文档解析未完成，请检查文件后重新上传。')
        finally:
            if process.poll() is None:
                process.kill()
            process.wait(timeout=3)
